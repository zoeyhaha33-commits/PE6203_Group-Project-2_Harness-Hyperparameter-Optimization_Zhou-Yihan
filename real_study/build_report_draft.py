"""Build an evidence-based English report draft with the bundled DOCX runtime."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'report_draft'
OUT.mkdir(exist_ok=True)
S = json.loads((ROOT / 'final_test_summary.json').read_text())
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.75)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in (('Title', 24), ('Heading 1', 16), ('Heading 2', 12)):
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(7)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('PE6203 Group Project 3 | Draft | ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
md = []
next_page = False
# The bundled default template contains paragraph borders in heading styles.
for st in doc.styles:
    if st.element.pPr is not None:
        for border in list(st.element.pPr.findall(qn('w:pBdr'))):
            st.element.pPr.remove(border)


def p(text):
    doc.add_paragraph(text)
    md.extend([text, ''])


def h(text, level=1):
    global next_page
    heading = doc.add_heading(text, level)
    if next_page:
        heading.paragraph_format.page_break_before = True
        next_page = False
    md.extend(['#' * (level + 1) + ' ' + text, ''])


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    t.autofit = False
    for c, text in zip(t.rows[0].cells, headers):
        c.text = str(text)
    pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader'); pr.append(repeat)
    for row in rows:
        for c, text in zip(t.add_row().cells, row):
            c.text = str(text)
    for row in t.rows:
        # Prevent one logical row splitting between pages.
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for i, c in enumerate(row.cells):
            if widths: c.width = Inches(widths[i])
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(4)
                para.paragraph_format.space_before = Pt(4)
                for run in para.runs: run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    md.extend(['| ' + ' | '.join(headers) + ' |', '|' + '|'.join(['---']*len(headers)) + '|'])
    md.extend('| ' + ' | '.join(map(str, row)) + ' |' for row in rows)
    md.append('')


def page():
    global next_page
    next_page = True


doc.add_heading('Harness Hyperparameter Optimization', 0)
md.extend(['# Harness Hyperparameter Optimization', ''])
p('A small scale empirical study of PI Agent on SkillsBench')
p('PE6203 Group Project 3 | Report draft | 30 October 2026')
p('Group members and student IDs: to be added by the authors.')
h('Abstract')
p('We study hyperparameter optimization over two configuration fields of a fixed PI Agent harness: contextWindow and maxTokens. We compare an untuned Default reference, Random search, Optuna TPE Bayesian optimization, and a student-developed two-stage multi-fidelity search. The three optimizers each receive nine validation task slots, with the selected configuration evaluated once on each of three disjoint held-out tasks. Search accounting totals USD 0.65175928 and held-out test accounting totals USD 0.31901342. Across 12 test slots, ten receive valid official rewards and all ten score zero; two slots have no valid reward. Bayesian optimization selects a candidate with one validation success out of three, but no test success is observed. The study demonstrates an auditable HPO workflow rather than a performance improvement. Sparse replication, repeated Bayesian configurations, unequal realized costs, and evolving failure-handling amendments limit comparisons. Exploratory matched-parameter analysis and raw-log inspection identify possible cost sensitivity and concrete output failures, without establishing causal parameter importance.')
if S.get('supplementary'):
    p('Update after a documented protocol amendment: one supplementary Scala evaluation for each originally missing Default and Bayesian slot produces valid reward 0 in both cases. All four supplemented three-task estimates are therefore 0%. The original ten scored outcomes and two missing outcomes remain the primary record; supplementary attempts and costs are reported separately in Appendix A.')
h('1 Research objective')
p('The objective is to select a supported harness configuration that performs well within a declared search budget and generalizes to unseen tasks. We ask whether systematic search improves task success, how quickly each method discovers its best complete validation candidate, and whether any gain warrants both search and subsequent execution costs. The optimization surface consists of configuration inputs; we do not optimize prompts, model weights, task definitions, or agent-loop source code.')
h('2 Experimental setting')
p('The fixed stack is PI coding agent 0.81.0 with pi-acp 0.0.34, orchestrated by BenchFlow 0.6.3 in Linux Docker containers. The model is litellm/qwen/qwen3-coder through OpenRouter. Runs use no-skill mode, a single-shot loop, and concurrency one. The benchmark repository is pinned to commit b63b7b2850226b6aa4fb5929a8c1ac7bc4d9a6af. These values come from the project protocol and installation records [1–3].')

page()
h('2 Experimental setting continued')
table(['Split', 'Tasks'], [
    ['Validation', 'financial-modeling-qa; jax-computing-basics; travel-planning'],
    ['Held-out test', 'python-scala-translation; reserves-at-risk-calc; threejs-structure-parser']], [1.0, 6.0])
p('The split contains three tasks per set with no task-name overlap. The 12 final test slots represent four methods on the same three tasks, not 12 independent benchmark questions. Task definitions and selected search states are fingerprinted in the frozen final protocol. Each final task has a 600-second agent limit; other task resource limits are recorded individually.')
table(['Field', 'Declared values'], [['contextWindow', '64000, 96000, 128000'], ['maxTokens', '4096, 8192, 16384']], [1.7, 5.3])
p('The Cartesian search space has nine possible configurations. The launcher passes these fields through the provider-model configuration supplied in BENCHFLOW_PROVIDER_MODELS. The Default reference supplies no parameter overrides; we therefore report it as the harness default rather than assign it unverified numeric values. This provider-model surface is narrower than unrestricted harness redesign. Metadata limits also should not be interpreted as proof of the exact number of tokens consumed in a run.')
h('3 Methods and selection rules')
p('Random search samples three unique configurations without replacement and evaluates each on all three validation tasks. Bayesian search uses Optuna TPE with one startup trial and three candidate evaluations. Duplicate configurations are allowed and treated as separate stochastic evaluations. In this run, all three Bayesian candidates use contextWindow=128000 and maxTokens=4096, so there is no evidence of broad parameter-space exploration by that method.')
p('Our developed method is a two-stage multi-fidelity search. It samples five unique configurations, evaluates each first on financial-modeling-qa, promotes two using the declared ranking, and evaluates their remaining two tasks. This uses five screening slots plus four continuation slots. Screening results are reused. The low-fidelity proxy reduces evaluation coverage per rejected candidate, but assumes that screening-task performance is informative about the remaining tasks. Here the screening task gives no successes, weakening that assumption.')
p('Only candidates with valid results on all three validation tasks may be selected. Ranking is by strict passes, then mean official reward, then lower accounted candidate cost, then earlier candidate index. The TPE objective itself is strict pass count. Final configurations are frozen before held-out evaluation; test results are not used to choose or retune them.')

page()
h('4 Budget and failure governance')
p('Each optimizer has nine validation slots and a USD 1.50 admission ceiling, with USD 0.20 reserved before admitting another task. The held-out phase shares a USD 2.00 ceiling. Admission checks are not hard provider spending caps. The protocol uses seed 20261028 for reproducible configuration generation; this does not make remote model responses deterministic.')
p('The realized histories are not identical: Random and multi-fidelity each have one additional recovery launch, for ten launches, while Bayesian has nine. Each of the first two methods also has one missing validation result. Equal admission ceilings do not require equal actual spending. The fairness limitation is asymmetric additional recovery opportunities, not unequal realized USD totals alone. We therefore describe this as a comparison under common declared slot and admission caps, with recovery deviations, rather than fully equalized actual attempts.')
p('Failed runs retain their records and budget intervals. A run without a valid reward is missing, not an imputed zero. During final testing, an agent reaching its frozen deadline with an official reward of zero is retained as a scored failure. A command timeout without a reward remains missing. Final task slots are not explicitly retried. Some failure classification and recovery rules were amended after observing failures, and timeout-specific reconciliation was implemented during final testing. Although configurations and task limits remained frozen, the treatment was not fully predeclared; this procedural deviation must be disclosed.')
p('Installation hardening was introduced during validation after the sixth multi-fidelity evaluation and applied to all final methods. It changes installer checks, network timeouts, and bounded installation retries while preserving pinned agent versions and agent logic. Consequently, validation runs do not all share exactly the same installation wrapper. Subsequent image preparation populated build caches without changing task Dockerfiles. These operational changes are documented rather than presented as HPO improvements.')
h('5 Validation search results')
rows=[]
for name, v in S['validation'].items():
    unique=len({tuple(sorted(c['config'].items())) for c in v['candidates']})
    rows.append([name, f"{len(v['candidates'])} / {unique}", f"{v['valid']} / {v['missing']}", v['launches'], f"{v['selected']['validation_pass_rate']:.1%}", f"{v['budget_accounted_usd']:.8f}"])
table(['Method', 'Candidates / unique', 'Scored / missing', 'Launches', 'Selected val.', 'Search USD'],rows,[1.25,1.15,1.1,.7,1.05,1.75])
p('Random selects candidate 0 with 64000/16384 and zero passes out of three. Its candidate 2 passes travel-planning but has a missing jax-computing-basics result and is ineligible. Bayesian selects candidate 1 with 128000/4096 and one success on jax-computing-basics. Multi-fidelity selects candidate 0 with 64000/16384 and zero passes out of three; its other promoted candidate is incomplete. Candidate indices are zero-based. Recovered real Default validation summaries show 1/3 success, on travel-planning. This historical baseline predates search and installation hardening; it is not synthetic demo data. It is examined separately in Section 9.')

page()
h('5 Validation search results continued')
doc.add_picture(str(ROOT/'validation_search_curves.png'), width=Inches(7.0))
md.extend(['![Validation search curves](../validation_search_curves.png)', ''])
p('Figure 1. Left: cumulative observed search accounting versus best-so-far validation pass rate among complete, eligible candidates. The curve is undefined until a complete candidate exists. Right: cumulative accounting by resolved validation slot. Both plots exclude held-out results. Equal slot caps do not imply equal realized USD spending.')
p('Bayesian first obtains a complete zero-pass candidate at slot 3, then reaches one-third validation pass rate at slot 6 and retains it. Random remains at zero among eligible candidates. Multi-fidelity obtains its first complete eligible candidate only at slot 9, so its incumbent is a single point rather than a meaningful early convergence trajectory. Its total observed cost is highest, despite the intended saving from screening. A single expensive screening evaluation contributes strongly to this result. The small sample does not show that multi-fidelity is generally inefficient.')
h('6 Primary held-out performance and generalization')
table(['Method', 'Frozen context / maxTokens', 'Test passes / scored', 'Missing', 'Full 3-task pass@1'],[
    ['Default','No overrides','0 / 2','1','Not identified'],
    ['Random','64000 / 16384','0 / 3','0','0.0%'],
    ['Bayesian','128000 / 4096','0 / 2','1','Not identified'],
    ['Multi-fidelity','64000 / 16384','0 / 3','0','0.0%']], [1.1,1.8,1.15,.65,2.3])
p('All ten scored final evaluations receive reward 0. For Default and Bayesian, 0/2 is the observed scored-only rate, not a fully observed three-task pass@1. With one unknown outcome, the logical full-set pass@1 bounds are 0% to 33.3%; these are missing-outcome bounds, not statistical confidence intervals. Default loses its Scala slot to a Docker build failure; Bayesian loses the same task to an unscored 900-second command timeout.')
p('The selected validation-to-test difference is 0 percentage points for Random and multi-fidelity. For Bayesian it is −33.3 percentage points when comparing 1/3 validation success with 0/2 observed test successes. This comparison is descriptive: the tasks differ and one Bayesian test result is missing. It is consistent with weak generalization but does not establish overfitting as the cause. Random and multi-fidelity select the same numeric configuration, so their final comparison cannot identify a parameter effect.')

page()
h('7 Search cost and runtime cost')
rows=[]
for name, m in S['methods'].items():
    v=S['validation'].get(name)
    rows.append([name, f"{v['budget_accounted_usd']:.8f}" if v else 'Not applicable',f"{m['budget_accounted_usd']:.8f}",f"{v['budget_accounted_usd']+m['budget_accounted_usd']:.8f}" if v else 'Not aggregated'])
table(['Method','Search USD','Held-out execution USD','Combined USD'],rows,[1.4,1.75,2.0,1.85])
p('The three searches total USD 0.65175928; the four test methods total USD 0.31901342. Their combined accounted amount is USD 0.97077270. This is not the complete project expenditure: earlier Default baseline work, provider probes, other diagnostics, and environment preparation are outside this aggregation. No search cost is assigned to the untuned Default reference.')
p('Costs are observed dedicated-key usage differences, subject to delayed billing. Failed intervals are preserved separately when attribution is uncertain. A stable zero difference is not proof that a run was free. Token telemetry is unavailable for some runs and is not substituted for USD cost. Because no method solves a held-out task, these data provide no demonstrated performance gain to amortize against optimization cost. A cost per solved task or a break-even deployment count would therefore be unsupported. Nor can the partial two-task execution totals for Default and Bayesian be treated as directly comparable full-workload deployment costs.')
h('8 Exploratory parameter sensitivity')
p('We conduct a post hoc sensitivity analysis using only existing validation records. Within the same task, we compare configurations that differ in one field while holding the other fixed. Repeated configurations are summarized by their mean observed cost and empirical success rate. The analysis contains 25 scored records, two missing records, 11 task/configuration cells, and six matched configuration pairs. It is an observational sensitivity method rather than a randomized ablation.')
table(['Task and fixed setting','Changed field','Low → high','n low / high','Success Δ','Mean USD Δ'],[
    ['Financial; context=64000','maxTokens','8192 → 16384','1 / 2','0 pp','+0.001203'],
    ['Financial; tokens=8192','contextWindow','64000 → 128000','1 / 1','0 pp','+0.127682'],
    ['Financial; tokens=4096','contextWindow','96000 → 128000','2 / 5','0 pp','+0.009938'],
    ['Financial; context=128000','maxTokens','4096 → 8192','5 / 1','0 pp','+0.143530'],
    ['JAX; tokens=4096','contextWindow','96000 → 128000','2 / 3','+33.3 pp','+0.001232'],
    ['Travel; tokens=4096','contextWindow','96000 → 128000','1 / 4','+25.0 pp','−0.050639']], [1.65,1.15,1.15,.7,.9,1.45])

page()
h('8 Exploratory parameter sensitivity continued')
p('All observed financial-task configurations fail, so changing either field does not produce an observed success improvement on that task. Some larger-context cells have higher empirical success on JAX and travel, but each effect is driven by very few successes. The largest cost contrast is observed when maxTokens increases from 4096 to 8192 at contextWindow=128000 on the financial task: mean observed cost rises by about USD 0.14353 without a success improvement. However, the high-value cell has only one run. The same expensive run also drives the large context-window contrast at maxTokens=8192, so these are not two independent pieces of evidence.')
p('The tentative finding is that the sampled context and output limits are associated with variable cost, while no reliable ranking of their influence on success is identifiable. Method, evaluation order, delayed billing, and model stochasticity are confounded. No causal claim about parameter importance or interaction is warranted. A stronger study would predeclare a replicated factorial design on validation tasks, randomize execution order, and use equal coverage per cell. Such a follow-up has not been run and would not change the frozen final results.')
page()
h('9 Failure mode comparison')
table(['Validation task','Default reward','Selected Bayesian reward','Observed contrast'],[
    ['financial-modeling-qa','0','0','Both fail'],
    ['jax-computing-basics','0','1','Default fails; selected candidate succeeds'],
    ['travel-planning','1','0','Default succeeds; selected candidate fails']], [1.7,1.0,1.5,2.8])
p('These two directional examples compare the actual selected Bayesian candidate 1 against the earlier untuned Default baseline on the same validation tasks. They supply outcome-level evidence requested by the failure-comparison criterion, but they are validation examples, not unseen test discoveries. The saved baseline summaries record clean wrapper exits, no agent/verifier errors, and official rewards. We copied them into validation_default_evidence/ and recorded hashes. Their original Linux jobs have not been independently re-read in this audit; specific baseline behavioral causes therefore remain unverified. Baseline task execution costs total USD 0.14043120, kept separate from HPO search costs and the previously reported cost scope.')
p('The contrast suggests that the selected configuration can trade success across tasks rather than dominate Default: JAX improves while travel deteriorates in these observations. However, model randomness, different run times and the later installation wrapper confound attribution. Bayesian candidates with identical numeric settings themselves vary between 0/3 and 1/3 validation successes. Thus the contrast does not prove that the parameter changes cause either improvement or harm, or establish overfitting. It also cannot isolate harmful interaction between contextWindow and maxTokens: the Default numeric values are not verified and there is no balanced factorial experiment. Interaction remains an untested hypothesis.')
p('For reserves-at-risk-calc, all four configurations pass only the no-errors-or-macros subtest and fail the three calculation stages plus the formula-presence check. The verifier finds zero formulas in required Gold price columns and missing or hardcoded contents in required Answer cells. Some numeric outputs are present, but they do not satisfy the required workbook computation structure. This is evidence of unmet output requirements, not an absence of execution [5].')
p('For threejs-structure-parser, all four runs pass output-directory existence and fail part/link mesh checks. Default and multi-fidelity omit cabin_N instance directories. Random uses a generic cabin directory; Bayesian uses cabins, neither matching the expected instances. Bayesian and multi-fidelity also produce a wheel_structure.obj without parsed vertices. These common structural failures are not repaired by the tested parameter changes [5].')
p('For python-scala-translation, Random reaches the 600-second agent limit after 13 tool calls but receives a valid official zero. Its verifier also finds a Double pattern incompatible with the expected Number type, causing source compilation failure. Multi-fidelity produces a zero with 21 test-compilation errors, including missing TokenizerBuilder and expected named arguments. Their heuristic code-quality score of 23/25 is not the official task reward and must not be reported as a 92% success rate. Default and Bayesian have no valid Scala reward and remain missing [5].')
p('Neither directional contrast exists on the final test set: all scored primary tasks and both supplements fail. The validation examples above must not be substituted for held-out evidence. Similar structural and compilation failures across test settings suggest a common bottleneck, but do not prove a unique root cause or exclude verifier defects. The old synthetic demo outputs are excluded from all these comparisons.')

page()
h('10 Reproducibility and evidence')
p('The repository should retain configuration generation, task launchers, frozen protocols, ledgers, raw logs, and derived analysis scripts. Commands must be run from the original WSL environment for live execution; the offline summaries can be regenerated without API calls. Credentials are supplied through environment variables and should not be committed. This draft does not claim that a public GitHub submission has already been created.')
table(['Artifact','Role'],[
    ['protocol.json and final_test_protocol.json','Search surface, split, budgets, pins, selected configurations and fingerprints'],
    ['run_search.py and run_test.py','Search scheduling, final launch guards, outcome and billing reconciliation'],
    ['run_test_batch.py and bench_hardened.py','Serial evaluation and installation-only hardening'],
    ['*_state.json and test_*_state.json','Validation histories, failed attempts, frozen test outcomes and costs'],
    ['final_test_logs.tar.gz and final_test_evidence/','Original test archive and indexed verifier evidence'],
    ['summarize_final_test.py and validation_summary.py','Tables, cost accounting, generalization comparison and curves'],
    ['analyze_parameter_sensitivity.py','Matched-parameter analysis using validation records only']], [2.9,4.1])
p('Offline reproduction from the project root: run python real_study/analyze_parameter_sensitivity.py, then python real_study/summarize_final_test.py. These scripts verify frozen search-state fingerprints and regenerate the analysis artifacts. Raw test results were checked against the 12 ledger task names and scored rewards, as well as available audit hashes. We did not rerun the verifier during the audit.')
h('11 Limitations and conclusion')
p('This is a small pilot with three validation tasks and three held-out tasks, not a statistically conclusive benchmark. The strongest threats are sparse replication, two missing final outcomes, unequal actual launch counts, post hoc failure-handling amendments, installation changes during validation, and limited independent parameter exploration. The held-out scores provide no evidence that any optimizer improves the harness over Default. Bayesian has the best observed eligible validation candidate but no observed held-out success; multi-fidelity does not achieve a realized cost saving in this run. The practical contribution is a reproducible, cost-accounted search workflow with explicit failure treatment and evidence-based negative results.')
p('Before submission, the group should add authorship, verify the exposed parameter pathway and repository instructions, review the documented fairness and protocol deviations, and decide whether the exploratory sensitivity evidence is adequate for the course or requires a separately declared validation follow-up. The completed held-out outcomes should remain unchanged.')
h('Evidence references', 2)
p('[1] Course brief, Group Project 3 Harness Hyperparameter Optimization, sections 3 and 4. [2] real_study/protocol.json and final_test_protocol.json. [3] bench_hardened.py and installation_hardening_check.json. [4] Validation/test ledgers and final_test_summary.json. [5] final_test_failure_analysis.md and final_test_evidence/manifest.json. [6] parameter_sensitivity.json and parameter_sensitivity.md. All experimental claims in this draft are derived from these project-local records; no external empirical results are asserted.')
if S.get('supplementary'):
    page()
    h('Appendix A Supplementary evaluations after protocol amendment')
    p('After reviewing the two unscored original failures, we recorded a separate amendment before additional model requests. It permits one supplementary python-scala-translation launch for Default and one for Bayesian. The original protocol, source ledgers and primary table are retained and fingerprinted. Default still supplies no parameter override, and Bayesian retains contextWindow=128000 and maxTokens=4096. The task definition, agent limits and pinned model/harness settings remain unchanged. Existing environment build caches may be reused. No previously scored zero is repeated and no best-of-attempt rule is applied.')
    extra=S['supplementary']
    table(['Method','Supplement status','Official reward','Extra observed USD'],
        [[name, r['status'], str(r['reward']), f"{r['cost_usd']:.8f}"] for name,r in extra['attempts'].items()], [1.4,1.9,1.4,2.3])
    p('Both supplements finish with return code 0 and official reward 0, with billing settled. Archived result hashes match the ledgers. Verifier logs show successful source compilation followed by 21 test-compilation errors in each run, including unresolved TokenType, Token and TokenizerBuilder symbols. No unit tests execute. Generated definitions appear in the trajectories, suggesting a package or test-integration mismatch, but this explanation has not been independently recompiled or proven. Heuristic code-quality scores of 88% for Default and 96% for Bayesian are not official task success rates. These scored failures supplement, rather than erase, the original missing outcomes.')
    table(['Method','Supplemented passes / tasks','Remaining missing','Supplemented pass@1'],
        [[name,f"{m['passes']} / {m['scored']}",m['missing'],f"{m['pass_rate']:.1%}" if m['pass_rate'] is not None else 'Not identified'] for name,m in extra['completed_estimates'].items()], [1.5,2.3,1.5,1.7])
    p(f"The two extra attempts cost USD {extra['additional_accounted_usd']:.8f}. Primary plus supplementary held-out accounting totals USD {extra['primary_plus_supplement_test_usd']:.8f}; including the three validation searches gives USD {extra['search_primary_and_supplement_usd']:.8f}. Earlier baseline work and diagnostic/environment expenses remain excluded. Supplementary attempts consume the original shared test admission budget rather than resetting it.")
    p('After supplementation, Bayesian has 0/3 held-out successes compared with 1/3 for its selected validation candidate, a descriptive gap of −33.3 percentage points. Random and multi-fidelity remain 0/3, and Default now has a supplemented estimate of 0/3. No observed test gain is demonstrated. The primary estimates with missing outcomes and the amended completed estimates should be displayed together, not substituted without qualification.')
    p('This amendment is post hoc and gives Default and Bayesian an extra launch opportunity. It is therefore a supplemental robustness/completion analysis, not a claim that all methods had equal numbers of attempts or that all failure rules were fixed before the original final evaluation. The amendment, two new ledgers, original failures and additional expenses must accompany the submission. No further retries are permitted by this amendment.')
    p('Supplementary evidence: supplementary_test/ ledgers, supplementary_test_logs.tar.gz and final_test_evidence/supplementary_manifest.json. No new configuration search was performed.')
page()
h('Appendix B Report requirement coverage and unresolved evidence')
p('This draft addresses all six report topics, but topic coverage is not the same as complete experimental support. The following audit replaces any earlier implication that the assignment was fully satisfied merely because test execution had ended.')
table(['Requirement','Evidence and location','Coverage judgment'],[
    ['Incumbent curves and equal budget','Sections 4–5; cumulative USD versus best complete-candidate validation score. Same declared 9 slots and USD 1.50 ceiling.','Qualified: equal caps do not require equal realized spending, but extra recovery launches were asymmetric.'],
    ['Final test pass@1','Section 6 and Appendix A; original missing outcomes retained; supplemented estimates all 0/3.','Reported, with post hoc supplementation explicitly labelled.'],
    ['Generalization gap','Section 6 and Appendix A; selected Bayesian 1/3 validation versus 0/3 supplemented test.','Reported descriptively; not causal proof of overfitting.'],
    ['Search versus runtime cost','Section 7 and Appendix A; original, search and supplement costs separated.','Reported; no demonstrated gain to justify optimization cost.'],
    ['Parameter sensitivity','Section 8; six task-matched one-parameter contrasts and cost deltas.','Partial: observations do not reliably rank parameter importance or isolate interactions.'],
    ['Failure-mode comparison','Section 9; Default/Bayesian success reversals on JAX and travel, plus held-out verifier cases.','Directional validation cases available; behavioral causes partly unverified and no held-out success reversals.']], [1.4,3.0,2.6])
p('A strict fairness comparison would require a prospective experiment with a common treatment of failed attempts, such as counting every wrapper launch against the same launch ceiling. Unequal actual USD totals alone are not a fairness violation when the declared caps are the same. Our asymmetric extra recovery opportunities, rather than cost inequality itself, are the key caveat. We do not retrospectively relabel the histories as fully equal-budget evidence.')
p('The main evidence gap is robust parameter sensitivity. A suitable follow-up would predeclare validation-only factorial cells, equal task coverage, replication, randomized run order, failure handling and a spending cap. Estimating harmful interactions requires changing both factors in a balanced design rather than inferring interactions from one-factor observations. No such follow-up has been executed or charged as part of this correction, and the held-out configurations and results remain frozen.')
p('Before submission, the group should review the evidence limits against the course rubric. This draft does not certify that a missing result or weak sensitivity estimate automatically satisfies the instructor. Additional experiments or clarification may be needed; additional successes cannot be guaranteed and should never be manufactured by repeated testing.')
doc.core_properties.title = 'Harness Hyperparameter Optimization'
doc.core_properties.subject = 'PE6203 Group Project 3 report draft'
doc.core_properties.author = ''
last = doc.paragraphs[-1]
last.paragraph_format.line_spacing = 1.0
last.paragraph_format.space_after = Pt(0)
for run in last.runs:
    run.font.size = Pt(9)
for para in doc.paragraphs:
    if para.text.startswith(('[1] Course brief', 'Supplementary evidence:')):
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_after = Pt(0)
        for run in para.runs:
            run.font.size = Pt(9)
doc.save(OUT/'PE6203_Harness_HPO_Report_Draft.docx')
(OUT/'PE6203_Harness_HPO_Report_Draft.md').write_text('\n'.join(md),encoding='utf-8')
print(OUT/'PE6203_Harness_HPO_Report_Draft.docx')
